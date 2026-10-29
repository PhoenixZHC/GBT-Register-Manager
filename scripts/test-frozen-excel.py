"""冻结组件的 Excel 保存验收；仅使用测试临时文件，不连接机器人。"""
import ctypes
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook

TARGET = Path(sys.argv.pop(1)).resolve()


def invoke(action, output, register_type='R', rows=None):
    result = subprocess.run([str(TARGET)], input=json.dumps(dict(action=action, registerType=register_type,
        rows=[dict(type='R', ID=1, value=42)] if rows is None else rows, outputPath=str(output))),
        capture_output=True, text=True, encoding='utf-8', timeout=30)
    if result.returncode:
        raise RuntimeError(result.stderr)
    return json.loads(result.stdout.split('<<<GBT-BEGIN>>>')[1].split('<<<GBT-END>>>')[0])


class FrozenExcelTests(unittest.TestCase):
    def test_mixed_pose_exports_and_templates(self):
        with tempfile.TemporaryDirectory() as folder:
            for kind in ('P', 'PR'):
                rows = json.loads((Path('.local-testing') / f'pose-{kind}.json').read_text(encoding='utf-8'))
                target = Path(folder) / f'{kind}.xlsx'
                self.assertTrue(invoke('export_excel', target, kind, rows)['ok'])
                wb = load_workbook(target, data_only=True)
                try:
                    values = list(wb.active.values)
                    self.assertEqual(values[0][2], 'PoseType')
                    for original, restored in zip(rows, values[1:]):
                        restored = dict(zip(values[0], restored))
                        for key, expected in original.items():
                            if isinstance(expected, float):
                                self.assertAlmostEqual(restored[key], expected, places=10)
                            else:
                                self.assertEqual(restored[key], expected)
                finally:
                    wb.close()
                self.assertTrue(invoke('export_template', target, kind, [])['ok'])
                wb = load_workbook(target)
                try:
                    self.assertEqual(wb.active.max_row, 1)
                    self.assertEqual(list(wb.active.values)[0], values[0])
                finally:
                    wb.close()

    def test_exports_replace_old_file_with_valid_workbook(self):
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / 'backup.xlsx'
            for action in ('export_excel', 'export_template'):
                with self.subTest(action=action):
                    target.write_bytes(b'old backup')
                    self.assertTrue(invoke(action, target)['ok'])
                    wb = load_workbook(target)
                    try:
                        expected = [('type', 'ID', 'value')]
                        if action == 'export_excel':
                            expected.append(('R', 1, 42))
                        self.assertEqual(list(wb.active.values), expected)
                    finally:
                        wb.close()

    @unittest.skipUnless(sys.platform == 'win32', 'Windows exclusive file lock')
    def test_locked_destination_is_preserved(self):
        kernel = ctypes.WinDLL('kernel32', use_last_error=True)
        kernel.CreateFileW.argtypes = [ctypes.c_wchar_p, ctypes.c_uint32, ctypes.c_uint32,
            ctypes.c_void_p, ctypes.c_uint32, ctypes.c_uint32, ctypes.c_void_p]
        kernel.CreateFileW.restype = ctypes.c_void_p
        kernel.CloseHandle.argtypes = [ctypes.c_void_p]
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / 'backup.xlsx'
            target.write_bytes(b'original backup')
            handle = kernel.CreateFileW(str(target), 0x80000000, 0, None, 3, 0x80, None)
            self.assertNotEqual(handle, ctypes.c_void_p(-1).value)
            try:
                for action in ('export_excel', 'export_template'):
                    with self.subTest(action=action):
                        self.assertFalse(invoke(action, target)['ok'])
            finally:
                kernel.CloseHandle(handle)
            self.assertEqual(target.read_bytes(), b'original backup')
            self.assertEqual(list(Path(folder).iterdir()), [target])


if __name__ == '__main__':
    unittest.main()
