"""真实 SDK 对象与 Excel 的往返；机器人接口全部为隔离替身。"""
import copy
import json
import math
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock
from openpyxl import Workbook, load_workbook
import test_register_safety as fixtures


class PoseRoundtripTests(unittest.TestCase):
    def setUp(self):
        fixtures.RegisterSafetyTests.setUp(self)
        fixtures.RegisterSafetyTests.pose_fixture(self)
        self.ns.update(Workbook=Workbook, Path=Path)

    def point(self, reg_type, kind, idx=1):
        pose = self.ns['ProgramPose']() if reg_type == 'P' else self.ns['PoseRegister']()
        pose.id = idx
        pose.name, pose.comment = '=point', '中文备注 =text'
        data = pose.poseData if reg_type == 'P' else pose.poseRegisterData
        data.pt = getattr(self.ns['PoseType'], kind)
        cart = data.cartData.baseCart if reg_type == 'P' else data.cartData
        for i, key in enumerate('xyzabc', 1):
            setattr(cart.position, key, i * 1.123456789)
        for i in range(1, 10):
            setattr(data.joint, f'j{i}', (-1)**i * i * 10.123456789)
        cart.posture.arm_left_right = 0
        cart.posture.wrist_flip = 1
        cart.posture.arm_up_down = -1
        cart.posture.arm_back_front = 0
        cart.posture.turn_cycle = [-4, -3, -2, -1, 0, 1, 2, 3, 4]
        if reg_type == 'P':
            data.cartData.tf, data.cartData.uf = 3, 4
        return pose

    def wire(self, pose, reg_type):
        if reg_type == 'PR':
            return pose.to_json()
        result = dict(pos=[], posture=dict(turnCircle=[]))
        pose.update_program_json(result)
        return result

    def test_mixed_excel_roundtrip_preserves_both_types_and_posture(self):
        for reg_type in ('P', 'PR'):
            poses = [self.point(reg_type, kind, i) for i, kind in enumerate(('CART', 'JOINT'), 1)]
            rows = [self.ns['pose_to_row'](p, reg_type, p.id) for p in poses]
            with tempfile.TemporaryDirectory() as folder:
                target = Path(folder) / 'mixed.xlsx'
                self.ns['export_excel'](reg_type, rows, str(target))
                wb = load_workbook(target, data_only=True)
                try:
                    values = list(wb.active.values)
                    imported = [dict(zip(values[0], values[i])) for i in range(1, 3)]
                finally:
                    wb.close()
            validated = self.ns['validate_rows'](reg_type, imported, 'overwrite')
            for source, row in zip(poses, validated):
                target = self.ns['ProgramPose']() if reg_type == 'P' else self.ns['PoseRegister']()
                target.id = source.id
                self.ns['apply_pose_row'](target, row, reg_type)
                actual = self.ns['pose_to_row'](target, reg_type, target.id)
                for key, expected in self.ns['pose_to_row'](source, reg_type, source.id).items():
                    if isinstance(expected, float):
                        self.assertAlmostEqual(actual[key], expected, places=11)
                    else:
                        self.assertEqual(actual[key], expected, key)
                self.assertEqual(self.wire(target, reg_type)['data']['pos_type'] if reg_type == 'PR'
                                 else self.wire(target, reg_type)['poseType'],
                                 source.poseRegisterData.pt.value if reg_type == 'PR' else row['PoseType'].lower())

    def test_write_switches_types_in_both_directions_and_skip_never_mutates(self):
        for reg_type in ('P', 'PR'):
            for old_kind, new_kind in (('CART', 'JOINT'), ('JOINT', 'CART')):
                original = self.point(reg_type, old_kind)
                source = self.point(reg_type, new_kind)
                row = self.ns['pose_to_row'](source, reg_type, 1)
                self.reg.read_PR.return_value = (original, 0)
                self.arm.program_pose.read_all_poses.return_value = ([original], 0)
                func = self.ns['write_p'] if reg_type == 'P' else self.ns['write_pr']
                args = (self.arm, 'test', row) if reg_type == 'P' else (self.arm, row)
                before = copy.deepcopy(self.wire(original, reg_type))
                self.assertEqual(func(*args, 'skip'), (True, 'skip'))
                self.assertEqual(self.wire(original, reg_type), before)
                self.assertEqual(func(*args, 'overwrite'), (True, 'write'))
                self.assertEqual(self.wire(original, reg_type), self.wire(source, reg_type))

    def test_joint_new_points_use_sdk_joint_payload(self):
        for reg_type in ('P', 'PR'):
            source = self.point(reg_type, 'JOINT')
            row = self.ns['pose_to_row'](source, reg_type, 1)
            self.reg.read_PR.return_value = (None, -28)
            self.arm.program_pose.read_all_poses.return_value = ([], 0)
            if reg_type == 'P':
                self.ns['write_p'](self.arm, 'test', row, 'overwrite')
                written = self.arm.program_pose.add.call_args.args[2]
                wire = self.wire(written, 'P')
                self.assertAlmostEqual(wire['pos'][8], math.radians(row['J9']))
            else:
                self.ns['write_pr'](self.arm, row, 'overwrite')
                written = self.reg.write_PR.call_args.args[0]
                self.assertEqual(written.to_json_new()['data']['pos']['data'][8], row['J9'])
            self.assertEqual(self.wire(written, reg_type), self.wire(source, reg_type))

    def test_bad_typed_rows_rejected_before_connect(self):
        row = self.ns['pose_to_row'](self.point('P', 'JOINT'), 'P', 1)
        for change in (dict(PoseType=''), dict(PoseType='other'), dict(J9=None), dict(J1='NaN'),
                       dict(X=1), dict(WristFlip=2), dict(Turn1=1.5), dict(TF=-1)):
            with self.subTest(change=change), self.assertRaisesRegex(RuntimeError, 'GBT_INVALID_REGISTER_DATA'):
                self.ns['apply_rows']({'controller_ip':'isolated'}, dict(registerType='P', programName='test',
                    rows=[dict(row, **change)], conflictPolicy='overwrite'))
        self.ns['connect_arm'].assert_not_called()

    def test_old_format_export_stays_compatible(self):
        row = dict(TYPE='PR', ID=1, X=1, Y=2, Z=3, A=4, B=5, C=6, coord='L')
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / 'old.xlsx'
            self.ns['export_excel']('PR', [row], str(target))
            wb = load_workbook(target)
            try:
                self.assertEqual(list(wb.active.values)[0], tuple(row))
            finally:
                wb.close()

    def test_frontend_fixture_export(self):
        folder = Path(__file__).resolve().parents[2] / '.local-testing'
        folder.mkdir(exist_ok=True)
        for reg_type in ('P', 'PR'):
            rows = [self.ns['pose_to_row'](self.point(reg_type, kind, i), reg_type, i)
                    for i, kind in enumerate(('CART', 'JOINT'), 1)]
            self.ns['export_excel'](reg_type, rows, str(folder / f'pose-{reg_type}.xlsx'))
            (folder / f'pose-{reg_type}.json').write_text(json.dumps(rows, ensure_ascii=False), encoding='utf-8')

    def test_new_sdk_pr_metadata_change_is_not_silently_ignored(self):
        self.reg.use_new_api = True
        row = self.ns['pose_to_row'](self.point('PR', 'JOINT'), 'PR', 1)
        self.reg.read_PR.return_value = (None, -28)
        for policy in ('skip', 'overwrite'):
            with self.assertRaisesRegex(RuntimeError, 'GBT_POSE_METADATA_UNSUPPORTED'):
                self.ns['apply_rows']({'controller_ip': 'isolated'}, dict(registerType='PR', rows=[row], conflictPolicy=policy))
        self.reg.write_PR.assert_not_called()
        row.update(Name='', Comment='')
        self.assertEqual(self.ns['write_pr'](self.arm, row, 'overwrite'), (True, 'write'))


if __name__ == '__main__':
    unittest.main()
